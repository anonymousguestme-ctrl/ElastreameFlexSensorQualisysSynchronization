from __future__ import annotations

from pathlib import Path
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\Elastreame电子皮肤\同步采集软件")
ASSET_DIR = ROOT / "manual_assets"
OUT = ROOT / "QTM电脑拷贝包_20260928" / "Qualisys与Elastreme同步采集图文说明书.docx"
QTM_SCREENSHOT = Path(r"C:\Users\anony\AppData\Local\Temp\codex-clipboard-6739d09e-5b0b-4d4a-9956-fb703d116eff.png")

NAVY = "17345E"
BLUE = "2563EB"
PALE_BLUE = "EAF2FF"
GREEN = "15966A"
AMBER = "D58A18"
RED = "D9475C"
INK = "172033"
MUTED = "667085"
LINE = "D9E1EC"
PALE = "F5F7FA"
WHITE = "FFFFFF"

FONT_REG = r"C:\Windows\Fonts\msyh.ttc"
FONT_BOLD = r"C:\Windows\Fonts\msyhbd.ttc"


def font(size: int, bold: bool = False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def round_box(draw, box, fill, outline=LINE, radius=24, width=3):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline="#" + outline, width=width)


def center_text(draw, box, text, fnt, fill="#172033", spacing=8):
    x1, y1, x2, y2 = box
    bb = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=spacing, align="center")
    w, h = bb[2] - bb[0], bb[3] - bb[1]
    draw.multiline_text(((x1+x2-w)/2, (y1+y2-h)/2), text, font=fnt, fill=fill, spacing=spacing, align="center")


def arrow(draw, start, end, color="#2563EB", width=8):
    draw.line([start, end], fill=color, width=width)
    ex, ey = end
    sx, sy = start
    import math
    angle = math.atan2(ey-sy, ex-sx)
    size = 20
    for delta in (2.55, -2.55):
        p = (ex + size*math.cos(angle+delta), ey + size*math.sin(angle+delta))
        draw.line([end, p], fill=color, width=width)


def save_flow():
    img = Image.new("RGB", (1600, 420), "white")
    d = ImageDraw.Draw(img)
    titles = ["插入接收器\n打开采集盒", "启动便携版\n选择 S1 S2", "进入待机\n确认有数据", "QTM 开始\n自动记录", "QTM 停止\n自动保存"]
    colors = ["#EAF2FF", "#EAF2FF", "#FFF5E3", "#FFECEF", "#EAF8F3"]
    x = 45
    for i, (title, c) in enumerate(zip(titles, colors), 1):
        box = (x, 105, x+255, 315)
        round_box(d, box, c)
        d.ellipse((x+18, 120, x+68, 170), fill="#2563EB")
        center_text(d, (x+18,120,x+68,170), str(i), font(25, True), "white")
        center_text(d, (x+30, 165, x+225, 285), title, font(31, True), spacing=12)
        if i < len(titles): arrow(d, (x+265,210), (x+315,210), "#7B91B3", 7)
        x += 315
    img.save(ASSET_DIR / "overall_flow.png", quality=95)


def save_topology():
    img = Image.new("RGB", (1600, 720), "white")
    d = ImageDraw.Draw(img)
    round_box(d, (535, 180, 1065, 545), "#EAF2FF", NAVY, 32, 4)
    center_text(d, (585, 210, 1015, 355), "运行 QTM 的电脑", font(47, True), "#17345E")
    center_text(d, (585, 355, 1015, 485), "QTM + 同步采集软件\nUDP 8989（默认）", font(31), "#324A6D", 12)
    round_box(d, (75, 120, 430, 320), "#EAF8F3", GREEN)
    center_text(d, (100, 145, 405, 295), "S1 传感器\n接收器 0051\n新电脑选择实际 COM", font(30, True), "#116B4D", 10)
    round_box(d, (75, 410, 430, 610), "#FFF5E3", AMBER)
    center_text(d, (100, 435, 405, 585), "S2 传感器\n接收器 0053\n新电脑选择实际 COM", font(30, True), "#87550E", 10)
    arrow(d, (430,220), (535,300), "#15966A")
    arrow(d, (430,510), (535,430), "#D58A18")
    round_box(d, (1170, 255, 1525, 470), "#F5F7FA", MUTED)
    center_text(d, (1200, 285, 1495, 440), "records\n每次采集生成\nCSV + 事件 + 元数据", font(29, True), "#39465B", 10)
    arrow(d, (1065,365), (1170,365), "#667085")
    img.save(ASSET_DIR / "topology.png", quality=95)


def save_ui_mock():
    img = Image.new("RGB", (1600, 980), "#F4F7FB")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((50, 45, 1550, 930), radius=28, fill="white", outline="#DDE4EE", width=3)
    d.rounded_rectangle((90, 85, 190, 175), radius=14, fill="#2563EB")
    center_text(d, (90,85,190,175), "Q × E", font(25, True), "white")
    d.text((220, 90), "Qualisys × Elastreme 同步采集", font=font(39, True), fill="#172033")
    d.text((220, 145), "QTM 开始时自动记录传感器数据", font=font(22), fill="#667085")
    labels = [("传感器", "2 / 2 有数据", GREEN), ("QTM 监听", "等待 QTM", GREEN), ("采集状态", "等待开始", AMBER)]
    x = 90
    for title, value, color in labels:
        round_box(d, (x, 210, x+450, 330), "#FBFCFE", LINE, 16, 2)
        d.ellipse((x+25,250,x+45,270), fill="#"+color)
        d.text((x+65, 232), title, font=font(20), fill="#667085")
        d.text((x+65, 270), value, font=font(27, True), fill="#172033")
        x += 480
    d.text((100, 382), "采集设置", font=font(30, True), fill="#172033")
    fields = [
        ("S1 · 0051", "COM8", "①"), ("S2 · 0053", "COM5", "②"),
        ("采样间隔", "20 ms", "③"), ("QTM 广播端口", "8989", "④"),
        ("保存位置", "...\\records", "⑤")
    ]
    y = 445
    for label, value, num in fields:
        d.text((115, y+15), label, font=font(23, True), fill="#2563EB")
        d.rounded_rectangle((420,y,960,y+64),radius=10,fill="#F8FAFC",outline="#DDE4EE",width=2)
        d.text((450,y+16),value,font=font(23),fill="#172033")
        d.ellipse((990,y+7,1040,y+57),fill="#2563EB")
        center_text(d,(990,y+7,1040,y+57),num,font(23,True),"white")
        y += 78
    d.rounded_rectangle((1110, 455, 1470, 535), radius=12, fill="#2563EB")
    center_text(d, (1110,455,1470,535), "进入待机", font(28, True), "white")
    d.rounded_rectangle((1110, 560, 1470, 640), radius=12, fill="#EEF2F7")
    center_text(d, (1110,560,1470,640), "停止监听", font(27, True), "#39465B")
    d.text((1110, 700), "⑥ 先看状态，再去 QTM 开始", font=font(25, True), fill="#D58A18")
    d.text((1110, 750), "绿色 2 / 2 有数据", font=font(23), fill="#15966A")
    d.text((1110, 790), "绿色 等待 QTM", font=font(23), fill="#15966A")
    img.save(ASSET_DIR / "ui_mock.png", quality=95)


def save_qtm_mock():
    img = Image.new("RGB", (1500, 600), "white")
    d = ImageDraw.Draw(img)
    round_box(d, (50,45,1450,550), "#F3F5F9", NAVY, 20, 3)
    d.rectangle((70,70,390,525), fill="#FFFFFF", outline="#C7D0DE", width=2)
    d.text((95,100), "Processing", font=font(30, True), fill="#172033")
    items=["Force Data","Gaze Vector","Real-Time Output","Euler Angles","TSV Export"]
    y=175
    for item in items:
        if item=="Real-Time Output":
            d.rounded_rectangle((88,y-8,370,y+42),radius=8,fill="#DDEAFF")
            d.text((105,y),item,font=font(24,True),fill="#174BAF")
        else:d.text((105,y),item,font=font(23),fill="#39465B")
        y+=62
    d.text((445,90), "Real-Time Output", font=font(35, True), fill="#172033")
    d.text((475,180), "☑  Use default port numbers", font=font(27), fill="#172033")
    d.text((475,260), "Capture Broadcast Port", font=font(27, True), fill="#172033")
    d.rounded_rectangle((970,235,1220,305),radius=10,fill="#FFFFFF",outline="#9AA8BC",width=2)
    center_text(d,(970,235,1220,305),"8989",font(30,True),"#17345E")
    d.text((475,365), "☐  Allow client control", font=font(27), fill="#172033")
    d.text((475,445), "端口必须与同步软件完全一致", font=font(26, True), fill="#D9475C")
    img.save(ASSET_DIR / "qtm_settings_mock.png", quality=95)


def save_capture_flow():
    img = Image.new("RGB", (1600, 400), "white")
    d = ImageDraw.Draw(img)
    steps=[("同步软件","进入待机"),("状态确认","2 / 2 有数据"),("QTM","Capture / Start"),("同步软件","正在记录"),("QTM","停止"),("同步软件","已保存 · 待机")]
    x=40
    for i,(a,b) in enumerate(steps):
        c="#EAF2FF" if i<3 else ("#FFECEF" if i==3 else "#EAF8F3")
        round_box(d,(x,95,x+220,300),c,LINE,18,2)
        center_text(d,(x+15,115,x+205,180),a,font(24,True),"#17345E")
        center_text(d,(x+15,185,x+205,275),b,font(25,True),"#172033")
        if i<len(steps)-1: arrow(d,(x+225,198),(x+272,198),"#7B91B3",6)
        x+=265
    img.save(ASSET_DIR / "capture_flow.png", quality=95)


def save_folder_tree():
    img=Image.new("RGB",(1400,600),"white")
    d=ImageDraw.Draw(img)
    d.text((80,55),"便携版文件夹",font=font(37,True),fill="#17345E")
    lines=[
        (0,"Qualisys_Elastreme_Sync.exe","启动软件"),
        (0,"config.json","保存通道与端口设置"),
        (0,"_internal","程序运行库 不能删除"),
        (0,"records","采集后自动生成"),
        (1,"试验名称_时间","每次 QTM 测量一个文件夹"),
        (2,"elastreme_raw.csv","S1 S2 原始数据与时间戳"),
        (2,"qtm_events.jsonl","QTM 开始停止消息"),
        (2,"session.json","配置 帧数 实测到帧率"),
    ]
    y=130
    for level,name,desc in lines:
        x=100+level*85
        color="#2563EB" if level==0 else ("#15966A" if level==1 else "#667085")
        d.text((x,y),"●",font=font(20),fill=color)
        d.text((x+38,y-2),name,font=font(25,True),fill="#172033")
        d.text((650,y),desc,font=font(23),fill="#667085")
        y+=55
    img.save(ASSET_DIR / "folder_tree.png",quality=95)


def save_troubleshoot():
    img=Image.new("RGB",(1600,620),"white")
    d=ImageDraw.Draw(img)
    round_box(d,(560,35,1040,150),"#EAF2FF",BLUE,20,3)
    center_text(d,(585,55,1015,130),"点击进入待机后是否 2 / 2 有数据",font(29,True),"#17345E")
    arrow(d,(650,150),(380,250),"#15966A",6); d.text((475,170),"否",font=font(23,True),fill="#D9475C")
    arrow(d,(950,150),(1220,250),"#15966A",6); d.text((1070,170),"是",font=font(23,True),fill="#15966A")
    round_box(d,(70,250,690,405),"#FFF5E3",AMBER,18,3)
    center_text(d,(95,270,665,385),"检查采集盒电源与配对\n关闭官方软件\n重新选择两个实际 COM 口",font(27,True),"#87550E",12)
    round_box(d,(910,250,1530,405),"#EAF8F3",GREEN,18,3)
    center_text(d,(935,270,1505,385),"QTM 开始后是否显示 正在记录",font(28,True),"#116B4D")
    arrow(d,(1220,405),(1100,495),"#D9475C",6)
    round_box(d,(650,485,1510,590),"#FFECEF",RED,18,3)
    center_text(d,(675,500,1485,575),"若没有 检查 QTM 与软件端口是否一致\n默认 8989 并检查 Windows 防火墙",font(26,True),"#8E2637",8)
    img.save(ASSET_DIR / "troubleshoot.png",quality=95)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LINE, size="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_alt_text(inline_shape, text):
    docPr = inline_shape._inline.docPr
    docPr.set("descr", text)
    docPr.set("title", text)


def add_figure(doc, path, caption, width=6.65):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    shape=p.add_run().add_picture(str(path),width=Inches(width))
    set_alt_text(shape, caption)
    p.paragraph_format.space_after=Pt(3)
    c=doc.add_paragraph()
    c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after=Pt(10)
    r=c.add_run(caption); set_run_font(r,size=9,color=MUTED)


def add_table(doc, headers, rows, widths=None):
    table=doc.add_table(rows=1,cols=len(headers))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    for i,h in enumerate(headers):
        cell=table.rows[0].cells[i]
        set_cell_shading(cell,NAVY); set_cell_border(cell)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(h);set_run_font(r,size=9.5,bold=True,color=WHITE)
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for i,val in enumerate(row):
            cell=cells[i];set_cell_border(cell);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: set_cell_shading(cell,"F7F9FC")
            p=cell.paragraphs[0]
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
            r=p.add_run(str(val));set_run_font(r,size=9.2)
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)
    return table


def add_bullet(doc, text, level=0):
    p=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2")
    p.paragraph_format.space_after=Pt(4)
    r=p.add_run(text);set_run_font(r)
    return p


def add_number(doc, lead, text):
    p=doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after=Pt(6)
    r=p.add_run(lead);set_run_font(r,bold=True)
    r=p.add_run(text);set_run_font(r)
    return p


def page_break(doc):
    doc.add_page_break()


def build_docx():
    doc=Document()
    sec=doc.sections[0]
    sec.page_height=Cm(29.7);sec.page_width=Cm(21)
    sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.55);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)

    styles=doc.styles
    normal=styles["Normal"]
    normal.font.name="Microsoft YaHei";normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
    normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing=1.2;normal.paragraph_format.space_after=Pt(6)
    title=styles["Title"];title.font.name="Microsoft YaHei";title._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
    title.font.size=Pt(28);title.font.bold=True;title.font.color.rgb=RGBColor(0,0,0)
    for style_name,size in [("Heading 1",19),("Heading 2",14)]:
        st=styles[style_name];st.font.name="Microsoft YaHei";st._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
        st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(7);st.paragraph_format.keep_with_next=True

    # Footer
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=fp.add_run("Qualisys 与 Elastreme 双通道同步采集图文说明书  |  2026-09-28")
    set_run_font(r,size=8,color=MUTED)

    # Cover
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(42);p.paragraph_format.space_after=Pt(16)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Qualisys 与 Elastreme\n双通道同步采集图文说明书")
    set_run_font(r,size=27,bold=True,color="000000")
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(20)
    r=p.add_run("QTM 电脑便携版操作与故障排查");set_run_font(r,size=14,color=MUTED)
    add_figure(doc,ASSET_DIR/"overall_flow.png","图 1  从连接设备到自动保存的完整流程",6.65)
    p=doc.add_paragraph()
    r=p.add_run("适用范围  ");set_run_font(r,bold=True)
    r=p.add_run("两个 Elastreme 高速单通道采集盒与 Qualisys QTM 的电脑端同步记录。S1 固定对应接收器 0051，S2 固定对应接收器 0053。");set_run_font(r)
    p=doc.add_paragraph()
    r=p.add_run("操作结论  ");set_run_font(r,bold=True)
    r=p.add_run("同步软件先进入待机并确认两个通道都有数据，再在 QTM 正常点击 Start。QTM 停止后，传感器文件会自动保存，无需同时手点两个软件。");set_run_font(r)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(18)
    r=p.add_run("版本 1.0    2026 年 9 月 28 日");set_run_font(r,size=10,color=MUTED)

    page_break(doc)
    doc.add_heading("1 系统组成与固定对应关系",level=1)
    p=doc.add_paragraph("便携版应放在运行 QTM 的 Windows 电脑上，两个 USB 接收器也插在这台电脑。这样 QTM 的 UDP 开始停止消息与串口到帧时间都由同一台电脑记录。")
    add_figure(doc,ASSET_DIR/"topology.png","图 2  QTM 电脑、两个接收器与记录文件的关系",6.6)
    add_table(doc,["通道","接收器编号","串口选择原则"],[
        ("S1","0051","按接收器外壳编号选择；换电脑后不一定仍是 COM8"),
        ("S2","0053","按接收器外壳编号选择；换电脑后不一定仍是 COM5"),
    ],[2.0,3.0,11.0])
    doc.add_heading("开始前检查",level=2)
    for t in [
        "两个采集盒均已打开，传感器连接牢固。",
        "0051、0053 两个 USB 接收器均插入 QTM 电脑。",
        "官方 Elastreme 软件已经关闭，避免占用串口。",
        "QTM 项目能够正常预览和采集。",
    ]: add_bullet(doc,t)

    page_break(doc)
    doc.add_heading("2 把便携版复制到 QTM 电脑",level=1)
    p=doc.add_paragraph("复制整个 QTM电脑拷贝包 文件夹。不能只复制 EXE，因为程序运行库和配置文件必须保留在同一文件夹内。")
    add_figure(doc,ASSET_DIR/"folder_tree.png","图 3  便携版目录与采集后生成的数据文件",6.4)
    add_number(doc,"复制完整文件夹。","将解压后的整个文件夹放到 QTM 电脑的本地磁盘，例如桌面或 D 盘实验目录。")
    add_number(doc,"启动程序。","双击 Qualisys_Elastreme_Sync.exe。目标电脑不需要安装 Python。")
    add_number(doc,"检查驱动。","如果下拉框没有出现 CH340 COM 口，运行 driver 文件夹中的 CH341SER.EXE，安装后重新插拔接收器。")
    add_number(doc,"保留目录结构。","不要删除 _internal 或 config.json。records 会在第一次进入待机时自动创建。")

    page_break(doc)
    doc.add_heading("3 软件界面与进入待机",level=1)
    add_figure(doc,ASSET_DIR/"ui_mock.png","图 4  同步采集软件界面与六个关键位置",6.65)
    add_table(doc,["编号","界面位置","正确设置"],[
        ("1","S1 · 0051","选择接收器 0051 在当前电脑对应的 COM 口"),
        ("2","S2 · 0053","选择接收器 0053 在当前电脑对应的 COM 口"),
        ("3","采样间隔","保持 20 ms"),
        ("4","QTM 广播端口","与 QTM Capture Broadcast Port 完全相同，默认 8989"),
        ("5","保存位置","建议保留便携版文件夹中的 records"),
        ("6","状态区","开始 QTM 前必须看到 2 / 2 有数据和等待 QTM"),
    ],[1.5,4.0,10.5])
    doc.add_heading("状态判断",level=2)
    p=doc.add_paragraph()
    r=p.add_run("2 / 2 已打开  ");set_run_font(r,bold=True,color=AMBER)
    r=p.add_run("只表示两个串口打开成功，还不能证明采集盒正在发送数据。");set_run_font(r)
    p=doc.add_paragraph()
    r=p.add_run("2 / 2 有数据  ");set_run_font(r,bold=True,color=GREEN)
    r=p.add_run("表示 S1、S2 均收到有效帧，可以开始实验。");set_run_font(r)

    page_break(doc)
    doc.add_heading("4 QTM 只需确认一次的设置",level=1)
    p=doc.add_paragraph("在 QTM 点击 Capture，打开 Start capture 窗口，再点击 Options。进入 Project Options 后，在左侧 Processing 分类中点击 Real-Time Output。")
    if QTM_SCREENSHOT.exists():
        add_figure(doc,QTM_SCREENSHOT,"图 5  QTM Project Options 左侧的 Real-Time Output 入口",5.05)
    p=doc.add_paragraph()
    r=p.add_run("定位方法  ");set_run_font(r,bold=True)
    r=p.add_run("左侧列表中 Real-Time Output 位于 Gaze Vector 下方、Euler Angles 上方。当前截图选中的是 Processing 总项，需要再点击 Real-Time Output。 ");set_run_font(r)
    doc.add_heading("4.1 核对 QTM 广播端口",level=2)
    add_figure(doc,ASSET_DIR/"qtm_settings_mock.png","图 6  Real-Time Output 页面需要核对的设置示意",4.9)
    add_number(doc,"确认端口。","Capture Broadcast Port 默认是 8989；同步软件中的 QTM 广播端口必须与它完全一致。")
    add_number(doc,"保持默认控制方式。","Allow client control 不需要勾选。我们的软件只是监听 QTM 的开始停止广播。")
    add_number(doc,"保存项目设置。","点击 OK。以后正常使用 QTM 的 Capture 和 Start 即可。")

    doc.add_heading("5 完成一次正式采集",level=1)
    add_figure(doc,ASSET_DIR/"capture_flow.png","图 7  一次采集中的软件状态变化",6.65)
    add_number(doc,"启动同步软件。","选择 S1、S2、采样间隔、QTM 广播端口和保存位置。")
    add_number(doc,"进入待机。","点击进入待机，等待传感器显示绿色 2 / 2 有数据，QTM 监听显示绿色 等待 QTM。")
    add_number(doc,"开始 QTM。","回到 QTM，设置测量名称和采集时间，点击 Start。同步软件应自动变为红色 正在记录。")
    add_number(doc,"完成动作。","采集期间不要拔接收器、关闭同步软件或启动官方 Elastreme 软件。")
    add_number(doc,"停止 QTM。","QTM 自动结束或手动停止后，同步软件自动显示 已保存 · 待机。")
    add_number(doc,"继续下一次。","直接在 QTM 开始下一次测量，不需要重新点击进入待机。全部实验完成后再点击停止监听。")
    doc.add_heading("第一次使用必须做的短测试",level=2)
    for t in [
        "做一次约 5 秒的 QTM 采集。",
        "确认同步软件从 等待开始 自动变为 正在记录。",
        "QTM 停止后确认 records 中出现新文件夹。",
        "打开 session.json，确认 S1、S2 都有帧数和实测到帧率。",
    ]: add_bullet(doc,t)

    page_break(doc)
    doc.add_heading("6 数据文件与后续对齐",level=1)
    p=doc.add_paragraph("每次 QTM 测量会在 records 中生成一个独立子文件夹，文件夹名称包含 QTM 测量名称和电脑时间。")
    add_table(doc,["文件","用途","重点检查"],[
        ("elastreme_raw.csv","S1、S2 原始值和电脑到帧时间","channel、relative_to_qtm_start_ms、frame_valid、pretrigger"),
        ("qtm_events.jsonl","QTM CaptureStart 和 CaptureStop 原始消息","开始和停止消息均存在"),
        ("session.json","通道、COM、端口、帧数和实测到帧率","S1、S2 均有 recorded_frames 和 receive rate"),
    ],[3.4,5.5,7.3])
    doc.add_heading("时间和采样率说明",level=2)
    p=doc.add_paragraph("CSV 使用长表保存，channel 列区分 S1 与 S2。两个串口各自记录到帧时间，后处理时再与 Qualisys 数据按时间戳对齐。")
    p=doc.add_paragraph("20 ms 是当前官方程序能够确认的最快串口请求，对应理论 50 Hz。手册中的 300 Hz 与当前程序行为存在冲突，因此应以 session.json 记录的实际到帧率为准。")
    p=doc.add_paragraph("当前方案属于电脑端 UDP 软件同步。时间零点是本机收到 QTM 开始广播的时刻，不应表述成硬件触发或固定毫秒精度同步。")

    page_break(doc)
    doc.add_heading("7 故障排查",level=1)
    add_figure(doc,ASSET_DIR/"troubleshoot.png","图 8  从串口状态到 QTM 触发的排查顺序",6.65)
    add_table(doc,["现象","原因方向","处理"],[
        ("找不到 COM 口","CH340 驱动或 USB 连接","安装 driver 中的驱动，重新插拔，点击刷新串口"),
        ("串口被占用","官方软件或其他串口工具正在使用","关闭其他程序，再重新进入待机"),
        ("2 / 2 已打开但没有数据","采集盒未上电或没有配对","检查两个采集盒电源、传感器连接和接收器对应关系"),
        ("只有 1 / 2 有数据","其中一路未传输或 COM 选错","按 0051、0053 外壳编号逐一核对"),
        ("QTM 开始但软件不记录","广播端口不同或防火墙拦截","确认两边端口完全一致；默认 8989；允许程序通过防火墙"),
        ("数据文件没有停止消息","QTM 异常退出或软件提前关闭","保留该次数据并重做短测试，正常试验中先停止 QTM"),
    ],[4.1,5.0,7.1])
    doc.add_heading("实验结束检查",level=2)
    for t in [
        "所有 QTM 测量均已停止并保存。",
        "同步软件当前没有显示 正在记录。",
        "records 中每次试验均有独立文件夹。",
        "最后点击停止监听，再关闭同步软件和采集盒。",
        "将 QTM 文件与对应的传感器记录文件夹一起备份。",
    ]: add_bullet(doc,t)

    OUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUT)
    print(OUT)


def main():
    ASSET_DIR.mkdir(parents=True,exist_ok=True)
    save_flow();save_topology();save_ui_mock();save_qtm_mock();save_capture_flow();save_folder_tree();save_troubleshoot()
    build_docx()


if __name__=="__main__": main()
